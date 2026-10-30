"""Spreadsheet parsing only. Google Sheets writes still use the original code."""
from pathlib import Path
from .protocol import JobError, main

MAX_FILE_BYTES = 12 * 1024 * 1024
MAX_ROWS = 50000
MAX_COLUMNS = 256
MAX_SHEETS = 64


class LimitedLogs(list):
    def append(self, value):
        if len(self) < 100:
            super().append(str(value)[:400])


def _legacy_parse_file(path, parser):
    import pandas as pd
    path = Path(path)
    if not path.is_file() or path.stat().st_size > MAX_FILE_BYTES:
        raise JobError("too_large")
    batting, pitching, logs = [], [], LimitedLogs()
    total_rows = 0

    def consume(frame, name):
        nonlocal total_rows
        total_rows += len(frame)
        if total_rows > MAX_ROWS or len(frame.columns) > MAX_COLUMNS:
            raise JobError("too_large")
        parser(frame, name, batting, pitching, logs)
        if len(batting) + len(pitching) > 10000:
            raise JobError("too_large")

    if path.suffix.lower() == ".csv":
        consume(pd.read_csv(path, encoding="utf-8-sig", header=None, nrows=MAX_ROWS + 1), "CSV_FILE")
    else:
        # A ZIP can expand to much more than its upload size.
        import zipfile
        if zipfile.is_zipfile(path):
            with zipfile.ZipFile(path) as archive:
                if sum(info.file_size for info in archive.infolist()) > 128 * 1024 * 1024:
                    raise JobError("too_large")
        with pd.ExcelFile(path) as workbook:
            if len(workbook.sheet_names) > MAX_SHEETS:
                raise JobError("too_large")
            for sheet in workbook.sheet_names:
                if sheet in ("설명서", "펑션"):
                    continue
                consume(workbook.parse(sheet_name=sheet, header=None, nrows=MAX_ROWS - total_rows + 1), sheet)
    return {"batting": batting, "pitching": pitching, "logs": logs}


def parse_file(path, parser=None):
    """Stream common formats; retain the legacy callback and optional engines.

    The production CSV/XLSX path never imports the Discord Cog, gspread or
    pandas. openpyxl may itself import numpy when it is installed.
    """
    path = Path(path)
    if not path.is_file() or path.stat().st_size > MAX_FILE_BYTES:
        raise JobError("too_large")
    if parser is not None:
        return _legacy_parse_file(path, parser)
    suffix = path.suffix.lower()
    if suffix in {".xls", ".xlsb", ".ods"}:
        # Keep previously supported engines/formats; still a disposable worker.
        from cogs.player_Record import PlayerRecord
        return _legacy_parse_file(path, PlayerRecord(None)._parse_single_sheet)
    if suffix not in {".csv", ".xlsx", ".xlsm"}:
        raise JobError("invalid_file")

    from .record_rows import parse_rows
    batting, pitching, logs = [], [], LimitedLogs()
    total_rows = 0

    def bounded(rows):
        nonlocal total_rows
        for row in rows:
            total_rows += 1
            if total_rows > MAX_ROWS or len(row) > MAX_COLUMNS:
                raise JobError("too_large")
            yield row
            if len(batting) + len(pitching) > 10000:
                raise JobError("too_large")

    if suffix == ".csv":
        import csv
        # Match the previous file cap and bound a single pathological CSV field.
        csv.field_size_limit(MAX_FILE_BYTES)
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = (row for row in csv.reader(stream) if row)
            parse_rows(bounded(rows), "CSV_FILE", batting, pitching, logs)
    else:
        import zipfile
        if not zipfile.is_zipfile(path):
            raise JobError("invalid_file")
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 4096 or sum(info.file_size for info in entries) > 128 * 1024 * 1024:
                raise JobError("too_large")
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            if len(workbook.sheetnames) > MAX_SHEETS:
                raise JobError("too_large")
            for sheet in workbook:
                if sheet.title in ("설명서", "펑션"):
                    continue
                # Reset potentially incorrect exported dimensions so data are
                # never silently ignored. Actual streamed rows are bounded.
                sheet.reset_dimensions()
                parse_rows(bounded(sheet.iter_rows(values_only=True)), sheet.title,
                           batting, pitching, logs)
        finally:
            workbook.close()
    return {"batting": batting, "pitching": pitching, "logs": logs}


def handle(request):
    return parse_file(request["path"])


if __name__ == "__main__":
    main(handle)
