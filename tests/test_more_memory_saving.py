"""v2 offline regressions. Real stub subprocesses; no account or AI model calls."""
from __future__ import annotations
import ast
import asyncio
import base64
import csv
import gc
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import types
import unittest
from unittest.mock import patch

import test_low_memory as previous
import resource_runtime as runtime
from resource_workers import excel_job
from resource_workers import record_rows

ROOT = Path(__file__).resolve().parents[1]
ROWS = [
    ['타자 기록', '', '', '', '', ''],
    ['선수명', '타수', '안타', '타점', '득점', '도루'],
    ['Batter[1]', 4, 2, 1, 2, 1],
    ['합계', '', '', '', '', ''],
    ['투수 기록', '', '', '', '', ''],
    ['선수명', '이닝', '타자', '피안타', '삼진', '자책점'],
    ['Pitcher (P)', 1.2, 7, 2, 2, 1],
    ['합계', '', '', '', '', ''],
]


class ConfigurationTests(unittest.TestCase):
    def test_disabled_message_cache_uses_none_not_zero(self):
        with patch.dict(os.environ, {'BOT_MESSAGE_CACHE': '0'}):
            self.assertIsNone(runtime.message_cache_limit())

    def test_message_cache_bound_and_invalid_value(self):
        for value, expected in [('20', 20), ('-1', None), ('100000', 1000), ('bad', None)]:
            with self.subTest(value=value), patch.dict(os.environ, {'BOT_MESSAGE_CACHE': value}):
                self.assertEqual(runtime.message_cache_limit(), expected)

    def test_no_cache_dependent_message_listeners(self):
        forbidden = {'on_message_edit', 'on_message_delete', 'on_bulk_message_delete',
                     'on_reaction_add', 'on_reaction_remove', 'on_reaction_clear'}
        for name in [ROOT / 'bot.py', *sorted((ROOT / 'cogs').glob('*.py'))]:
            tree = ast.parse(name.read_text())
            names = {n.name for n in ast.walk(tree) if isinstance(n, ast.AsyncFunctionDef)}
            self.assertFalse(names & forbidden, (name, names & forbidden))

    def test_python_311_syntax(self):
        for name in [ROOT / 'bot.py', ROOT / 'resource_runtime.py', ROOT / 'cogs/automod.py',
                     *sorted((ROOT / 'resource_workers').glob('*.py'))]:
            ast.parse(name.read_text(), filename=str(name), feature_version=(3, 11))

    def test_cpu_models_and_member_features_not_removed(self):
        self.assertIn('INTENTS.members = True', (ROOT / 'bot.py').read_text())
        self.assertIn('download.pytorch.org/whl/cpu', (ROOT / 'Dockerfile').read_text())
        self.assertIn('AUTOMOD_OCR_IDLE_SECONDS=0', (ROOT / 'Dockerfile').read_text())


class GateTests(unittest.IsolatedAsyncioTestCase):
    async def test_one_owner_and_fair_order(self):
        gate = runtime.HeavyWorkerGate()
        first = await gate.acquire('first')
        order = []
        async def run(name):
            lease = await gate.acquire(name)
            try:
                self.assertEqual(gate.owner.label, name)
                order.append(name)
                await asyncio.sleep(.01)
            finally:
                lease.release()
        tasks = [asyncio.create_task(run(str(i))) for i in range(3)]
        await asyncio.sleep(.01)
        self.assertEqual(order, []); self.assertEqual(gate.waiters, 3)
        first.release()
        await asyncio.wait_for(asyncio.gather(*tasks), 2)
        self.assertEqual(order, ['0', '1', '2']); self.assertIsNone(gate.owner)

    async def test_waiter_wakes_owner(self):
        event = asyncio.Event(); gate = runtime.HeavyWorkerGate()
        first = await gate.acquire('first', on_waiter=event.set)
        task = asyncio.create_task(gate.acquire('second'))
        await asyncio.wait_for(event.wait(), 1)
        first.release(); lease = await task; lease.release()

    async def test_cancelled_waiter_cannot_release_another_owner(self):
        gate = runtime.HeavyWorkerGate(); owner = await gate.acquire('owner')
        task = asyncio.create_task(gate.acquire('cancelled'))
        await asyncio.sleep(.01); task.cancel()
        with self.assertRaises(asyncio.CancelledError): await task
        self.assertIs(gate.owner, owner); self.assertEqual(gate.waiters, 0)
        owner.release(); new = await gate.acquire('new'); new.release()

    async def test_timeout_and_double_release_are_safe(self):
        gate = runtime.HeavyWorkerGate(); first = await gate.acquire('first')
        with self.assertRaises(TimeoutError): await gate.acquire('timeout', timeout=.01)
        first.release(); second = await gate.acquire('second'); first.release()
        self.assertIs(gate.owner, second); second.release(); self.assertEqual(gate.waiters, 0)

    async def test_per_loop_gate_is_shared(self):
        self.assertIs(runtime.get_heavy_gate(), runtime.get_heavy_gate())


class CombinedLifecycleTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        worker = self.root / 'ocr.py'; worker.write_text(previous.OCR_STUB)
        self.ns = previous.extract_ocr(worker, self.root)
        self.pool = self.ns['OCRPool'](self.root, idle_seconds=0, timeout_seconds=1)
        self.manager = runtime.WorkManager()
        d = self.root / 'resource_workers'; d.mkdir(); (d / '__init__.py').write_text('')
        (d / 'pdf_job.py').write_text(previous.JOB_STUB)
        self.patcher = patch.object(runtime, 'ROOT', self.root); self.patcher.start()

    async def asyncTearDown(self):
        await self.pool.close(); await self.manager.close()
        self.patcher.stop(); self.tmp.cleanup()
        self.assertIsNone(runtime.get_heavy_gate().owner)
        self.assertEqual(runtime.get_heavy_gate().waiters, 0)

    async def test_zero_idle_releases_before_result_and_restarts(self):
        await self.pool.start(eager=False)
        for _ in range(4):
            result = await self.pool.submit(b'img', 'ok')
            self.assertEqual(result.status, 'ok'); self.assertEqual(self.pool.readiness, 'standby')
            self.assertIsNone(self.pool._proc); self.assertIsNone(runtime.get_heavy_gate().owner)

    async def test_document_waits_for_real_ocr_exit(self):
        self.pool.idle_seconds = -1
        await self.pool.start(); pid = self.pool.health['worker_pid']
        jobtask = asyncio.create_task(self.pool.submit(b'img', 'slow'))
        await asyncio.sleep(.04)
        async with self.manager.slot('pdf') as job:
            result = await asyncio.wait_for(job.run('pdf_job', {}), 3)
            self.assertFalse(Path(f'/proc/{pid}').exists())
        self.assertEqual(result['answer'], 42); self.assertEqual((await jobtask).status, 'ok')

    async def test_persistent_idle_ocr_yields_to_document(self):
        self.pool.idle_seconds = -1; await self.pool.start()
        pid = self.pool.health['worker_pid']
        async with self.manager.slot('pdf') as job:
            self.assertEqual((await asyncio.wait_for(job.run('pdf_job', {}), 3))['answer'], 42)
        self.assertFalse(Path(f'/proc/{pid}').exists()); self.assertEqual(self.pool.readiness, 'standby')

    async def test_ocr_waits_without_spawning_while_document_holds_slot(self):
        lease = await runtime.get_heavy_gate().acquire('pdf')
        await self.pool.start(eager=False)
        task = asyncio.create_task(self.pool.submit(b'img', 'ok'))
        await asyncio.sleep(.05)
        self.assertIsNone(self.pool._proc); self.assertFalse(task.done())
        lease.release(); self.assertEqual((await asyncio.wait_for(task, 2)).status, 'ok')

    async def test_ocr_close_cancels_waiter_without_leaking_slot(self):
        lease = await runtime.get_heavy_gate().acquire('pdf')
        await self.pool.start(eager=False)
        task = asyncio.create_task(self.pool.submit(b'img', 'ok'))
        await asyncio.sleep(.05); await self.pool.close()
        self.assertEqual((await task).status, 'closed')
        self.assertIs(runtime.get_heavy_gate().owner, lease); lease.release()

    async def test_document_close_cancels_slot_waiter_before_spawn(self):
        lease = await runtime.get_heavy_gate().acquire('fake-ocr')
        async def run():
            async with self.manager.slot('pdf') as job: await job.run('pdf_job', {})
        task = asyncio.create_task(run()); await asyncio.sleep(.05)
        await self.manager.close()
        with self.assertRaises(runtime.WorkerError): await asyncio.wait_for(task, 1)
        self.assertEqual(self.manager._processes, set()); lease.release()

    async def test_crash_releases_gate_for_document(self):
        await self.pool.start(eager=False)
        self.assertEqual((await self.pool.submit(b'img', 'crash')).status, 'error')
        async with self.manager.slot('pdf') as job:
            self.assertEqual((await job.run('pdf_job', {}))['answer'], 42)

    async def test_wait_timeout_not_claimed_as_ocr_success(self):
        lease = await runtime.get_heavy_gate().acquire('pdf')
        self.pool.startup_timeout_seconds = .03
        await self.pool.start(eager=False)
        result = await asyncio.wait_for(self.pool.submit(b'img', 'ok'), 1)
        self.assertEqual(result.status, 'unavailable'); self.assertIsNone(self.pool._proc)
        lease.release()

    async def test_message_cache_and_gate_visible_in_diagnostic(self):
        bot = types.SimpleNamespace(get_cog=lambda _: types.SimpleNamespace(ocr=self.pool))
        report = runtime.snapshot_text(bot)
        self.assertIn(runtime.VERSION, report); self.assertIn('공유 작업 슬롯:', report)


class ProtocolTests(unittest.IsolatedAsyncioTestCase):
    async def test_chunked_encoding_is_identical_and_writes_are_small(self):
        with tempfile.TemporaryDirectory() as d:
            ns = previous.extract_ocr(Path(d) / 'unused', Path(d))
            pool = ns['OCRPool'](d)
            class Writer:
                def __init__(self): self.parts = []; self.drains = 0
                def write(self, data): self.parts.append(bytes(data))
                async def drain(self): self.drains += 1
            class Reader:
                async def readline(self): return b'{"status":"no_text","lines":[]}\n'
            for size in (1, 2, 3, 49151, 49152, 49153, 8*1024*1024):
                with self.subTest(size=size):
                    writer = Writer(); pool._proc = types.SimpleNamespace(stdin=writer, stdout=Reader())
                    data = bytes((i % 251 for i in range(size)))
                    job = ns['_Job'](data, '한글".png', 'image/png', '', 0, asyncio.get_running_loop().create_future())
                    await pool._exchange(job)
                    request = json.loads(b''.join(writer.parts))
                    self.assertEqual(base64.b64decode(request['data'], validate=True), data)
                    self.assertEqual(request['filename'], '한글".png')
                    self.assertLessEqual(max(map(len, writer.parts)), 65536)
            pool._proc = None

    async def test_sleeping_image_setting_validates_model_before_enable(self):
        tree = ast.parse((ROOT/'cogs/automod.py').read_text())
        cls = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == 'AutoMod')
        method = next(n for n in cls.body if getattr(n, 'name', None) == 'image_setting')
        method.decorator_list = []
        module = ast.Module(body=[ast.ImportFrom(module='__future__', names=[ast.alias(name='annotations')], level=0), method], type_ignores=[])
        changes = []; replies = []
        async def respond(*args): replies.append(args)
        ns = {'respond': respond}; exec(compile(ast.fix_missing_locations(module), '<setting>', 'exec'), ns)
        class FakeOCR:
            ready = False
            called = 0
            result = True
            async def start(self): self.called += 1; return self.result
        async def change(*args, **kwargs): changes.append(kwargs)
        obj = types.SimpleNamespace(ocr=FakeOCR(), _change=change)
        await ns['image_setting'](obj, None, '켜기')
        self.assertEqual(obj.ocr.called, 1); self.assertEqual(changes, [{'image_enabled': True}])
        obj.ocr.result = False
        with self.assertRaises(ValueError): await ns['image_setting'](obj, None, '켜기')
        self.assertEqual(len(changes), 1)
        await ns['image_setting'](obj, None, '끄기')
        self.assertEqual(changes[-1], {'image_enabled': False})


class StreamingRecordsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name); self.ns = previous.extract_parser()

    def write_csv(self, rows=ROWS):
        path = self.root / 'sample.csv'
        with path.open('w', encoding='utf-8-sig', newline='') as stream: csv.writer(stream).writerows(rows)
        return path

    def compare(self, path):
        expected = excel_job.parse_file(path, lambda *a: self.ns['_parse_single_sheet'](None, *a))
        result = excel_job.parse_file(path)
        for key in ('batting', 'pitching'): self.assertEqual(result[key], expected[key])
        return result

    def test_csv_matches_previous_parser(self): self.compare(self.write_csv())

    def test_xlsx_matches_previous_parser_and_closes_file(self):
        from openpyxl import Workbook
        book = Workbook(); sheet = book.active; sheet.title = '경기'
        for row in ROWS: sheet.append(row)
        book.create_sheet('설명서').append(['invalid content to skip'])
        path = self.root / 'sample.xlsx'; book.save(path); book.close()
        self.compare(path); path.unlink(); self.assertFalse(path.exists())

    def test_missing_values_and_nonfinite_values_match_previous(self):
        rows = [ROWS[0], ROWS[1]]
        for i, value in enumerate(['', 'nan', 'inf', '-inf', 'NULL', '=1+2', '1,234', '-2', '0.3']):
            rows.append([f'B{i}', value, value, value, value, value])
        rows.extend(ROWS[3:]); self.compare(self.write_csv(rows))

    def test_section_and_name_cleanup_matches_previous(self):
        rows = [ROWS[0], ROWS[1], ['A(test)[2]', '4', '2', '', '', ''],
                ['123', '4', '', '', '', ''], ['선수명', '', '', '', '', ''],
                ['합계', '', '', '', '', ''], ['Ignored', '9', '', '', '', ''],
                *ROWS[4:], ['IgnoredAfterTotal', '5', '', '', '', '']]
        self.compare(self.write_csv(rows))

    def test_production_csv_worker_imports_no_ai_dataframe_or_sdk(self):
        path = self.write_csv()
        code = "import sys,json; from resource_workers.excel_job import handle; r=handle({'path':sys.argv[1]}); assert r['batting'][0]['안타']==2; assert not any(x in sys.modules for x in ('pandas','numpy','discord','gspread','torch')); print('OK')"
        result = subprocess.run([sys.executable, '-S', '-c', code, str(path)], cwd=ROOT, capture_output=True, text=True, timeout=10)
        self.assertEqual(result.returncode, 0, result.stderr)

    def test_production_xlsx_worker_has_no_pandas_or_discord(self):
        from openpyxl import Workbook
        book=Workbook()
        for row in ROWS: book.active.append(row)
        path=self.root/'sample.xlsx';book.save(path);book.close()
        # The test host's sitecustomize preloads pandas. -S plus an explicit
        # dependency directory isolates application imports from that harness.
        import openpyxl
        packages = str(Path(openpyxl.__file__).resolve().parents[1])
        code="import sys; sys.path.insert(0,sys.argv[2]); from resource_workers.excel_job import handle; r=handle({'path':sys.argv[1]}); assert r['batting'][0]['안타']==2; assert not any(x in sys.modules for x in ('pandas','discord','gspread','torch'))"
        r=subprocess.run([sys.executable,'-S','-c',code,str(path),packages],cwd=ROOT,capture_output=True,text=True,timeout=20)
        self.assertEqual(r.returncode,0,r.stderr)

    def test_row_column_and_record_limits(self):
        path = self.write_csv()
        with patch.object(excel_job, 'MAX_ROWS', 2), self.assertRaises(excel_job.JobError): excel_job.parse_file(path)
        with patch.object(excel_job, 'MAX_COLUMNS', 2), self.assertRaises(excel_job.JobError): excel_job.parse_file(path)
        path = self.write_csv([ROWS[0], ROWS[1]] + [[f'B{i}', 1, 0, 0, 0, 0] for i in range(10001)])
        with self.assertRaises(excel_job.JobError): excel_job.parse_file(path)

    def test_sheet_and_zip_limits(self):
        from openpyxl import Workbook
        book=Workbook();book.create_sheet('extra');path=self.root/'many.xlsx';book.save(path);book.close()
        with patch.object(excel_job,'MAX_SHEETS',1),self.assertRaises(excel_job.JobError):excel_job.parse_file(path)

    def test_invalid_file_is_not_silently_empty(self):
        path=self.root/'bad.xlsx';path.write_text('not a zip')
        with self.assertRaises(excel_job.JobError):excel_job.parse_file(path)

    def test_fake_sheet_dimensions_cannot_hide_rows(self):
        from openpyxl import Workbook
        from zipfile import ZipFile, ZIP_DEFLATED
        import re
        book=Workbook()
        for row in ROWS:book.active.append(row)
        path=self.root/'dimensions.xlsx';book.save(path);book.close()
        with ZipFile(path) as z:entries={i.filename:z.read(i) for i in z.infolist()}
        key='xl/worksheets/sheet1.xml'
        entries[key]=re.sub(rb'<dimension ref="[^"]+"',b'<dimension ref="A1:A1"',entries[key])
        with ZipFile(path,'w',ZIP_DEFLATED) as z:
            for name,data in entries.items():z.writestr(name,data)
        result=excel_job.parse_file(path)
        self.assertEqual(result['batting'][0]['선수명'],'Batter')
        self.assertEqual(result['pitching'][0]['선수명'],'Pitcher')

    def test_debug_output_bounded_for_stream(self):
        path=self.write_csv([ROWS[0],ROWS[1]]+[[f'B{i}',1,0,0,0,0] for i in range(500)])
        result=excel_job.parse_file(path)
        self.assertLessEqual(len(result['logs']),100)
        self.assertTrue(all(len(s)<=400 for s in result['logs']))


if __name__ == '__main__': unittest.main(verbosity=2)
